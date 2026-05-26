#!/usr/bin/env python3
"""
先輩診断士の活動報告_20260215.pptx を Aisunia テンプレートに変換するスクリプト
使用方法: python migrate_to_aisunia.py
"""
import sys
import io
import zipfile
import copy
from io import BytesIO
from lxml import etree

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

# ===== パス設定 =====
OLD_PATH = r"C:\Users\Panasonic\OneDrive\เอกสาร\PPT\提案書\先輩診断士の活動報告_20260215.pptx"
POTX_PATH = r"C:\Users\Panasonic\Documents\Projects\homepage-builder\brand\slides\outputs\Aisunia_Template.potx"
TMP_TEMPLATE = r"C:\Users\Panasonic\AppData\Local\Temp\aisunia_migrate_tmp.pptx"
OUTPUT_PATH = r"C:\Users\Panasonic\OneDrive\เอกสาร\PPT\提案書\先輩診断士の活動報告_20260512.pptx"

# ===== potx → pptx 変換 =====
print("テンプレート変換中...")
with zipfile.ZipFile(POTX_PATH, "r") as zin:
    with zipfile.ZipFile(TMP_TEMPLATE, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.template.main+xml",
                    b"presentationml.presentation.main+xml",
                )
            zout.writestr(item, data)

old_prs = Presentation(OLD_PATH)
new_prs = Presentation(TMP_TEMPLATE)

# テンプレートに含まれるデフォルトスライドを削除
while len(new_prs.slides) > 0:
    sldIdLst = new_prs.slides._sldIdLst
    sldId = sldIdLst[0]
    rId = sldId.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    sldIdLst.remove(sldId)
    if rId in new_prs.part.rels:
        new_prs.part.rels.pop(rId)

master = new_prs.slide_masters[0]
LAYOUTS = {layout.name: layout for layout in master.slide_layouts}
print(f"旧スライド数: {len(old_prs.slides)}")
print(f"利用可能レイアウト: {list(LAYOUTS.keys())}")

# ===== セクション区切りレイアウト内テキストボックス位置 =====
# （レイアウトの「01」「セクションタイトル」を上書きするテキストボックスの位置）
SEC_NUM_L, SEC_NUM_T, SEC_NUM_W, SEC_NUM_H = (
    Inches(0.701), Inches(1.799), Inches(4.0), Inches(2.0),
)
SEC_TITLE_L, SEC_TITLE_T, SEC_TITLE_W, SEC_TITLE_H = (
    Inches(5.965), Inches(2.0), Inches(7.0), Inches(1.201),
)

# ===== セクションタイトル → 番号・表示テキスト マッピング =====
SECTION_MAP = {
    "自己紹介": ("01", "自己紹介"),
    "やってよかったこと": ("02", "やってよかったこと"),
    "やっておいたほうがよかったこと": ("03", "やっておいたほうが\nよかったこと"),
    "これからの展望": ("04", "これからの展望"),
}


# ===== ヘルパー =====
def get_slide_title(slide):
    for shape in slide.shapes:
        if shape.is_placeholder:
            pt = shape.placeholder_format.type.name
            if pt in ("TITLE", "CENTER_TITLE"):
                # \x0b はPowerPointの「ソフト改行」→ スペースに変換
                return shape.text_frame.text.replace("\x0b", " ").strip()
    return ""


def set_title(slide, text):
    if not text:
        return
    for shape in slide.shapes:
        if shape.is_placeholder:
            pt = shape.placeholder_format.type.name
            if pt in ("TITLE", "CENTER_TITLE"):
                tf = shape.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = text
                return


def add_textbox(slide, left, top, width, height, text,
                font_size=None, bold=False, rgb=None, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    if font_size:
        run.font.size = Pt(font_size)
    if bold:
        run.font.bold = True
    if rgb:
        run.font.color.rgb = RGBColor(*rgb)
    return box


def copy_image_to_slide(src_slide, dst_slide, old_rId):
    """画像を src_slide から dst_slide にコピーし新 rId を返す"""
    if old_rId not in src_slide.part.rels:
        return None
    rel = src_slide.part.rels[old_rId]
    if "image" not in rel.reltype.lower():
        return None
    img_blob = rel.target_part.blob
    try:
        _, new_rId = dst_slide.part.get_or_add_image_part(BytesIO(img_blob))
        return new_rId
    except Exception as e:
        print(f"    [警告] 画像コピー失敗 rId={old_rId}: {e}")
        return None


def copy_user_shapes(src_slide, dst_slide):
    """プレースホルダー以外のシェイプを src → dst にコピー"""
    src_spTree = src_slide.shapes._spTree
    dst_spTree = dst_slide.shapes._spTree
    count = 0

    for elem in list(src_spTree):
        tag = etree.QName(elem.tag).localname

        if tag == "sp":
            # プレースホルダーはスキップ
            nvSpPr = elem.find(qn("p:nvSpPr"))
            if nvSpPr is not None:
                nvPr = nvSpPr.find(qn("p:nvPr"))
                if nvPr is not None and nvPr.find(qn("p:ph")) is not None:
                    continue
            dst_spTree.append(copy.deepcopy(elem))
            count += 1

        elif tag == "pic":
            elem_copy = copy.deepcopy(elem)
            for blip in elem_copy.findall(".//" + qn("a:blip")):
                old_rId = blip.get(qn("r:embed"))
                if old_rId:
                    new_rId = copy_image_to_slide(src_slide, dst_slide, old_rId)
                    if new_rId:
                        blip.set(qn("r:embed"), new_rId)
            dst_spTree.append(elem_copy)
            count += 1

        elif tag in ("graphicFrame", "grpSp", "cxnSp"):
            dst_spTree.append(copy.deepcopy(elem))
            count += 1

    return count


def determine_layout(idx, slide):
    old_layout = slide.slide_layout.name
    title = get_slide_title(slide)

    if idx == 0:
        return "表紙", "COVER"
    if old_layout == "目次":
        return "目次", "TOC"
    if old_layout == "ユーザー設定レイアウト":
        if title in SECTION_MAP:
            return "セクション区切り", "SECTION"
        elif title == "質疑応答" or idx >= 25:
            return "エンド", "END"
        else:
            return "セクション区切り", "SECTION"
    return "コンテンツ", "CONTENT"


# ===== メイン変換 =====
print("\nスライド変換開始...")

for i, old_slide in enumerate(old_prs.slides):
    layout_name, layout_type = determine_layout(i, old_slide)
    new_slide = new_prs.slides.add_slide(LAYOUTS[layout_name])
    title = get_slide_title(old_slide)

    if layout_type == "COVER":
        set_title(new_slide, title or "アクセルパートナーズ　実務実習先輩診断士の活動報告")
        print(f"  [{i+1:2d}] 表紙: {title}")

    elif layout_type == "TOC":
        n = copy_user_shapes(old_slide, new_slide)
        print(f"  [{i+1:2d}] 目次: {n}個のシェイプをコピー")

    elif layout_type == "SECTION":
        sec_num, sec_title_text = SECTION_MAP.get(title, ("--", title or "セクション"))
        # レイアウトの固定テキストをスライドレベルのテキストボックスで上書き
        add_textbox(
            new_slide,
            SEC_NUM_L, SEC_NUM_T, SEC_NUM_W, SEC_NUM_H,
            sec_num, font_size=72, bold=True, rgb=(255, 255, 255),
            align=PP_ALIGN.CENTER,
        )
        add_textbox(
            new_slide,
            SEC_TITLE_L, SEC_TITLE_T, SEC_TITLE_W, SEC_TITLE_H,
            sec_title_text, font_size=28, bold=True, rgb=(26, 43, 74),
        )
        print(f"  [{i+1:2d}] セクション区切り: {sec_num} {sec_title_text!r}")

    elif layout_type == "END":
        print(f"  [{i+1:2d}] エンド: {title or '(エンドスライド)'}")

    else:  # CONTENT
        set_title(new_slide, title)
        n = copy_user_shapes(old_slide, new_slide)
        print(f"  [{i+1:2d}] コンテンツ: 「{title[:25] if title else 'タイトルなし'}」{n}個のシェイプ")

# ===== 保存 =====
print(f"\n保存中: {OUTPUT_PATH}")
new_prs.save(OUTPUT_PATH)
print(f"✅ 完了！スライド数: {len(new_prs.slides)}")
